"""Reviewed public electricity profiles and Torino municipal context.

Raw workbooks and PVGIS responses remain in _raw. No individual bills,
meters, or company records are required or read by this pipeline.
"""
import datetime as dt
import hashlib
import io
import json
from pathlib import Path
import zipfile

import openpyxl
import requests

ROOT = Path(__file__).resolve().parents[1]
DATE = "2026-10-22"
ARERA = "https://www.arera.it/dati-e-statistiche/dettaglio/analisi-dei-consumi-dei-clienti-domestici"
ISTAT = "https://www.istat.it/comunicato-stampa/trasformazioni-e-innovazioni-nei-consumi-elettrici-delle-citta-capoluogo-anno-2024/"
ISTAT_FILE = "https://www.istat.it/wp-content/uploads/2026/07/ENERGIA_serie_storiche-1.xlsx"


def write(name, data):
    (ROOT / "data" / name).write_text(json.dumps(data, ensure_ascii=False, separators=(",", ":")), encoding="utf-8")


def cached(path, url, params=None):
    target = ROOT / "_raw" / path
    target.parent.mkdir(parents=True, exist_ok=True)
    if not target.exists():
        response = requests.get(url, params=params, timeout=120)
        response.raise_for_status()
        target.write_bytes(response.content)
        target.with_suffix(target.suffix + ".source.json").write_text(json.dumps({"url": response.url, "retrieved": DATE, "sha256": hashlib.sha256(response.content).hexdigest()}), encoding="utf-8")
    return target


def main():
    archive = ROOT / "_raw/arera/Dati_regione_2025.zip"
    z = zipfile.ZipFile(archive)
    member = next(n for n in z.namelist() if "Dati prelievo clienti" in n)
    book = openpyxl.load_workbook(io.BytesIO(z.read(member)), read_only=True, data_only=True)
    cohorts = {}
    for r in book.active.values:
        if r[4] != "Piemonte" or r[1] != "Tutti" or r[3] != "Residente":
            continue
        cohort = cohorts.setdefault(r[2], {"power_class": r[2], "monthly": [], "annual_source": None})
        if isinstance(r[0], dt.datetime):
            cohort["monthly"].append({"month": r[0].month, "kwh": r[8]})
        elif str(r[0]) == "2025":
            assert cohort["annual_source"] is None
            cohort["annual_source"] = r[8]
    for c in cohorts.values():
        c["monthly"].sort(key=lambda r: r["month"])
        assert [r["month"] for r in c["monthly"]] == list(range(1, 13))
        c["annual_kwh"] = sum(r["kwh"] for r in c["monthly"])
        assert abs(c["annual_kwh"] - c["annual_source"]) < 1e-6
    # Hourly customer means have a different population from monthly means.
    # Preserve them as reference profiles; never assert simultaneous consumption.
    hourly_member = next(n for n in z.namelist() if "regionale1_5_a_3" in n)
    hourly = openpyxl.load_workbook(io.BytesIO(z.read(hourly_member)), read_only=True, data_only=True)
    profiles = []
    keys = set()
    for r in hourly.active.values:
        if r[1:5] != ("Piemonte", "Tutti", "1.5<potenza_impegnata<=3", "Residente"):
            continue
        month = r[0].month if isinstance(r[0], dt.datetime) else 0
        hour = int(str(r[6]).removeprefix("Ora"))
        key = (month, r[5], hour)
        assert key not in keys and 1 <= hour <= 24 and r[7] >= 0
        keys.add(key)
        profiles.append({"month": month, "day_type": r[5], "source_hour": hour, "kwh": r[7]})
    assert len(profiles) == 13 * 3 * 24
    write("public_demand_profiles.json", {
        "provider": "ARERA", "year": 2025, "region": "Piemonte", "residency": "Residente", "market": "Tutti", "retrieved": DATE,
        "source": ARERA, "data_url": "https://www.arera.it/fileadmin/allegati/dati/monitoraggioretail/Dati_regione_2025.zip",
        "raw_sha256": hashlib.sha256(archive.read_bytes()).hexdigest(), "monthly_member": member, "hourly_member": hourly_member,
        "cohorts": list(cohorts.values()), "hourly_power_class": "1.5<potenza_impegnata<=3", "hourly_reference": profiles,
        "monthly_semantics": "Sum of changing-population monthly mean withdrawals, not a fixed-cohort annual bill. Supply-to-household mapping is an explicit scenario assumption.",
        "hourly_semantics": "Means for hourly-treated customers only; source hour labels Ora1..Ora24 retained without inferred UTC/DST alignment. Not paired with PV output for self-consumption calculation.",
    })
    path = cached("istat/ENERGIA_serie_storiche-1.xlsx", ISTAT_FILE)
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    metrics = []
    for sheet, key, unit in [
        ("Tav.2.1", "domestic_electricity", "toe/100 residents"),
        ("Tav.2.2", "domestic_electricity_total", "toe"),
        ("Tav.6.2", "pv_installations", "count"),
        ("Tav.6.5", "pv_installed_capacity", "MW"),
        ("Tav.6.7", "pv_net_generation", "GWh"),
        ("Tav.6.10", "pv_self_consumed_generation", "GWh"),
    ]:
        rows = list(wb[sheet].values)
        headers = next(r for r in rows if 2024 in r)
        candidates = [r for r in rows if str(r[0]).split(" (")[0] == "Torino"]
        assert len(candidates) == 1
        torino = candidates[0]
        series = [{"year": year, "value": torino[i]} for i, year in enumerate(headers) if isinstance(year, int) and 2000 <= year <= 2024]
        assert series[-1]["year"] == 2024 and all(isinstance(r["value"], (int, float)) and r["value"] >= 0 for r in series)
        metrics.append({"key": key, "sheet": sheet, "source_title": rows[0][0], "source_row": torino[0], "notes": [r[0] for r in rows if isinstance(r[0], str) and (r[0].startswith("(") or r[0].startswith("Fonte"))], "unit": unit, "series": series})
    write("torino_energy_context.json", {"provider": "ISTAT", "municipality": "Torino", "municipality_code": "001272", "published": "2026-07-29", "retrieved": DATE,
        "source": ISTAT, "data_url": ISTAT_FILE, "raw_sha256": hashlib.sha256(path.read_bytes()).hexdigest(), "metrics": metrics,
        "scope": "Whole municipality, all installations/covered uses in each table, not San Donato, individual roofs or households. Context only; not spatially allocated into property estimates."})
    # A separate meteorological-year scenario, kept distinct from multi-year PVcalc means.
    scenarios = {}
    for name, angle, aspect in [("south", 30, 0), ("east", 30, -90), ("west", 30, 90), ("flat", 0, 0)]:
        params = dict(lat=45.078, lon=7.654, peakpower=1, loss=14, angle=angle, aspect=aspect, pvcalculation=1,
                      startyear=2023, endyear=2023, raddatabase="PVGIS-SARAH3", outputformat="json")
        path = cached(f"pvgis/hourly-2023-{name}.json", "https://re.jrc.ec.europa.eu/api/v5_3/seriescalc", params)
        data = json.loads(path.read_text(encoding="utf-8"))
        rows = data["outputs"]["hourly"]
        assert len(rows) == 8760 and len({r["time"] for r in rows}) == 8760
        times = [dt.datetime.strptime(r["time"], "%Y%m%d:%H%M") for r in rows]
        assert all((b-a).total_seconds() == 3600 for a, b in zip(times, times[1:]))
        assert all(r["P"] >= 0 for r in rows)
        monthly = [{"month": m, "kwh_per_kwp": sum(r["P"] / 1000 for r, t in zip(rows, times) if t.month == m)} for m in range(1, 13)]
        scenarios[name] = {"inputs": data["inputs"], "source_params": params, "source": requests.Request("GET", "https://re.jrc.ec.europa.eu/api/v5_3/seriescalc", params=params).prepare().url,
                           "raw_sha256": hashlib.sha256(path.read_bytes()).hexdigest(), "monthly": monthly,
                           "annual_kwh_per_kwp": sum(r["kwh_per_kwp"] for r in monthly), "samples": len(rows),
                           "first_timestamp": rows[0]["time"], "last_timestamp": rows[-1]["time"]}
        print(name, scenarios[name]["annual_kwh_per_kwp"], flush=True)
    write("solar_weather_2023.json", {"provider": "European Commission JRC PVGIS 5.3", "retrieved": DATE, "year": 2023,
        "semantics": "Hourly modelled power P in W, integrated using 1 h steps /1000 into kWh per 1 kWp; monthly bins use source timestamps. One weather year, not a multi-year mean or urban roof survey.", "scenarios": scenarios})
    print(f"Prepared {len(cohorts)} demand cohorts, {len(profiles)} hourly reference bins, {len(metrics)} municipal series and four weather scenarios.")


if __name__ == "__main__":
    main()
